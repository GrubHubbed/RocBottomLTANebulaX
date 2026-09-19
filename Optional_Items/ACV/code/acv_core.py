"""ACV Thermal Integrity Depot Console -- ranking engine.

Physically-motivated heuristic that ranks the 8 cars of a trainset from most to
least likely to have a refrigerant leak.  No Streamlit import lives in this
module: app.py and predict.py both consume it, so there is exactly one scoring
code path and the console and the submission CLI cannot disagree.

Nothing here reads Train_Labels.csv.  Ground truth is used by validate.py only.
"""

from __future__ import annotations

import io
import math
import os
import re
import time
from collections import Counter, OrderedDict
from dataclasses import dataclass, field

import openpyxl

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------

# Single gate for the two ranking bases.  See RANKING BASES below.
RANKING_PREFER_ASYMMETRY = True

# Two cars whose shortfalls differ by less than this are thermally
# indistinguishable and are separated by the cross-car z-score instead.
TIE_BAND_C = 0.05

# A filter that would leave fewer rows than this is the wrong filter for the
# schema in front of us, so it is dropped rather than applied.
MIN_ROWS_ABS = 20
MIN_ROWS_FRAC = 0.05

# Sniffing budget.  Phase 1 never reads more than this many data rows.
SNIFF_ROWS = 5
SNIFF_ROWS_ESCALATED = 60

# A channel needs this many paired observations on both circuits to be used.
MIN_CIRCUIT_ROWS = 20

# Literal strings seen in the workbooks where a blank cell was meant.
# acv_case_02 writes 'None' into empty cells; without this, numeric columns
# look categorical and the file silently produces eight "no data" cars.
NULL_LIKE = {"none", "nan", "null", "na", "n/a", "-", "--", "?", ""}

CAR_HEADER_RE = re.compile(r"^\s*Car\s*(\d+)\s*-\s*(.+?)\s*$", re.IGNORECASE)

# Severity is measured against the trainset's OWN fleet, not against an absolute
# temperature.  Absolute bands calibrated on car model A flagged all eight cars
# on acv_case_06 (model C, early July, fleet-mean setpoint 22.5 C, the coldest in
# the dataset): every unit sits near capacity, the fleet median shortfall is
# +0.92 C, and a fixed threshold therefore indicts the whole train.
#
# Median and median absolute deviation are used rather than mean and standard
# deviation so the one genuinely faulty car cannot drag the baseline toward
# itself.  The transform is monotonic WITHIN a file, so the within-file order is
# mathematically identical to the raw metric -- which is why it cannot affect any
# ranking.  Ranking continues to sort on raw shortfall / raw circuit asymmetry.
Z_BANDS = [(8.0, "CRITICAL"), (4.0, "HIGH"), (2.0, "ELEVATED")]

SEVERITY_ACTION = {
    "CRITICAL": "Withdraw at next depot entry; pressure-test the circuit",
    "HIGH": "Schedule refrigerant check within 48 h",
    "ELEVATED": "Watchlist; re-read after next service day",
    "NORMAL": "No action",
    "NO DATA": "Check the data feed before judging",
}

# The MAD floor stops a tightly matched fleet from producing absurd z values.
# It is per-metric because the two ranking metrics carry different units: a
# shortfall is degrees Celsius, an asymmetry is a dimensionless ratio.  Applying
# the temperature floor to the ratio would swamp it -- on acv_case_04 the real
# MAD is 0.011, so a 0.05 floor would pull car 01 from z 22.3 down to 4.8 and
# under-report a leak that is 8.2x clear of the next car.
MAD_FLOOR_C = 0.05          # cooling shortfall, degrees Celsius
MAD_FLOOR_RATIO = 0.005     # circuit asymmetry, dimensionless

# A fleet whose median car is running this far above its own setpoint is telling
# you about the weather or the setpoint, not about eight separate faults.
FLEET_SHORTFALL_ADVISORY_C = 0.50
SEVERITY_ORDER = ["CRITICAL", "HIGH", "ELEVATED", "NORMAL", "NO DATA"]
SEVERITY_ICON = {
    "CRITICAL": "!!!",
    "HIGH": "!!",
    "ELEVATED": "!",
    "NORMAL": "OK",
    "NO DATA": "X",
}

# A lead of at least this much is reported as clear rather than narrow.
CLEAR_LEAD_C = 0.30
NARROW_LEAD_C = 0.10

# The same question on the asymmetry basis is a ratio, not a temperature: car 01
# on acv_case_04 leads by 0.237 in absolute terms but by 8.2x, which is
# emphatically a clear lead.  Judging it against a Celsius threshold would
# mislabel it "narrow".
PRIME_RATIO_CLEAR = 2.00
PRIME_RATIO_NARROW = 1.25

# Confidence on the cross-car z-score margin.
MARGIN_CONFIDENT = 1.25
MARGIN_PROBABLE = 1.08

# The cross-car z-score earns its power from breadth -- it is the method that
# finds the pressure columns in the 63-parameter schema.  Averaged over a single
# parameter it is not an independent signal at all: in the 8-parameter schema the
# only measured, non-shared parameter left is the cabin temperature, which is the
# numerator of the shortfall itself.  Below this many parameters the z-score is
# reported as corroboration but is not allowed to reorder anything.
MIN_Z_PARAMS_FOR_TIEBREAK = 2

# A dropout difference smaller than this fraction of the recording is incidental
# rather than evidence, and does not break a tie.  Matches the 5% materiality
# used by the sensor-health rules.
MATERIAL_DROPOUT_FRAC = 0.05


# --------------------------------------------------------------------------
# Role table
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class RoleSpec:
    name: str
    required: bool
    numeric: bool           # role must land on a numeric column
    aliases: tuple          # exact header bases, in priority order
    keywords: tuple         # substring patterns, in priority order
    negative: tuple         # substrings that disqualify a candidate outright


ROLES = (
    RoleSpec(
        "outcome", True, True,
        ("indoor average temperature", "passenger cabin temperature detected value"),
        ("indoor average temperature", "passenger cabin temperature",
         "cabin temperature", "indoor temperature"),
        # 'Observation Area Temperature' is a different measurement point.
        ("observation area", "outdoor", "outside", "fresh air", "target",
         "control temperature", "setpoint", "set point", "return air"),
    ),
    RoleSpec(
        "reference", True, True,
        ("acv control temperature (cooling)", "target temperature value"),
        ("control temperature (cooling)", "target temperature value",
         "cooling setpoint", "cooling set temperature"),
        # Reject the offset-selection flags Target Temperature -2K/-1K/0/+1K/+2K.
        # They are booleans, not temperatures.
        ("heating", "+1k", "+2k", "-1k", "-2k", "target temperature 0",
         "target temperature +", "target temperature -"),
    ),
    # Not a ranking input.  Section 4E rejected the heating setpoint as a
    # feature -- it is flat for every car -- but it is read so that the
    # sensor-health rules can see a function that is switched off for the
    # whole service, which is a finding in its own right.
    RoleSpec(
        "heating_reference", False, True,
        ("acv control temperature (heating)",),
        ("control temperature (heating)", "heating setpoint"),
        ("cooling",),
    ),
    RoleSpec(
        "running_mode", False, False,
        ("acv running mode", "acv operating mode"),
        ("running mode", "operating mode"),
        (),
    ),
    RoleSpec(
        "information_valid", False, False,
        ("acv information valid",),
        ("information valid",),
        (),
    ),
    RoleSpec(
        "outdoor_temperature", False, True,
        ("outdoor average temperature", "outside temperature sensor reading",
         "fresh air temperature detected value"),
        ("outdoor average temperature", "outside temperature",
         "fresh air temperature", "ambient temperature"),
        ("valve", "damper"),
    ),
    RoleSpec(
        "load_derate", False, False,
        ("load halved", "load shedding"),
        ("load halved", "load shedding"),
        (),
    ),
    RoleSpec(
        "low_pressure", False, True,
        ("refrigeration system 1 low pressure value",),
        ("low pressure",),
        ("fault", "alarm", "switch"),
    ),
    RoleSpec(
        "high_pressure", False, True,
        ("refrigeration system 1 high pressure value",),
        ("high pressure",),
        ("fault", "alarm", "switch"),
    ),
    RoleSpec(
        "compressor_running", False, True,
        ("compressor 1 running",),
        ("compressor running", "compressor 1 running"),
        ("fault", "cab ", "auxiliary"),
    ),
    RoleSpec(
        "setting_mode", False, False,
        ("acv setting mode", "acv control mode"),
        ("setting mode", "control mode"),
        (),
    ),
)

ROLE_BY_NAME = {r.name: r for r in ROLES}
REQUIRED_ROLES = tuple(r.name for r in ROLES if r.required)

# Per-circuit channels, detected by regex on the header base, independently of
# the single-instance role table above.
CIRCUIT_PATTERNS = OrderedDict((
    ("lp", re.compile(r"refrigeration system (\d+).*low pressure", re.IGNORECASE)),
    ("hp", re.compile(r"refrigeration system (\d+).*high pressure", re.IGNORECASE)),
    ("duty", re.compile(r"^compressor (\d+) running$", re.IGNORECASE)),
    ("sv", re.compile(r"refrigeration system (\d+).*solenoid valve open", re.IGNORECASE)),
))
CIRCUIT_LABEL = {
    "lp": "low pressure",
    "hp": "high pressure",
    "duty": "compressor duty",
    "sv": "solenoid valve",
    "ratio": "compression ratio",
}


# --------------------------------------------------------------------------
# Value helpers
# --------------------------------------------------------------------------

def is_null_like(value) -> bool:
    """True for a real blank and for the literal strings used in place of one."""
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    if isinstance(value, str):
        return value.strip().lower() in NULL_LIKE
    return False


def to_number(value):
    """Coerce to float, or None.  Booleans count as 0/1; null-likes as None."""
    if is_null_like(value):
        return None
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        text = value.strip().replace(",", "")
        try:
            return float(text)
        except ValueError:
            return None
    return None


def to_text(value):
    """Lower-cased text, or None for a null-like."""
    if is_null_like(value):
        return None
    return str(value).strip().lower()


def nan_mean(values):
    live = [v for v in values if v is not None]
    return sum(live) / len(live) if live else None


def classify_nature(samples):
    """numeric / categorical / empty, from a small sample of cell values."""
    live = [v for v in samples if not is_null_like(v)]
    if not live:
        return "empty"
    numeric = sum(1 for v in live if to_number(v) is not None)
    return "numeric" if numeric >= max(1, math.ceil(0.6 * len(live))) else "categorical"


def car_key(car) -> str:
    """Zero-padded two-digit id, taken verbatim from the file's own headers."""
    digits = re.sub(r"\D", "", str(car))
    return digits.zfill(2) if digits else str(car)


# --------------------------------------------------------------------------
# Phase 1 -- sniff
# --------------------------------------------------------------------------

@dataclass
class Column:
    index: int
    header: str
    car: object
    param: str
    nature: str = "empty"
    role: object = None
    matched_by: object = None
    circuit_channel: object = None
    circuit_index: object = None


@dataclass
class Schema:
    sheet_title: str
    n_rows: int
    n_cols: int
    columns: list
    cars: list
    roles: dict                # car -> role -> Column
    circuits: dict             # car -> channel -> {circuit index -> Column}
    meta: dict                 # 'car model' / 'train number' / 'time' -> Column
    sniff_rows_used: int
    escalated: bool
    log: list = field(default_factory=list)

    def role_column(self, car, role):
        return self.roles.get(car, {}).get(role)


def _open_largest_sheet(source):
    """Open the workbook read-only and pick the largest sheet from metadata.

    Sheet choice uses ws.max_row * ws.max_column, which openpyxl knows from the
    worksheet dimension record -- the sheet is never loaded to decide this.
    """
    if isinstance(source, (bytes, bytearray)):
        handle = io.BytesIO(source)
    else:
        handle = source
    workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
    sheet = max(
        workbook.worksheets,
        key=lambda ws: (ws.max_row or 0) * (ws.max_column or 0),
    )
    return workbook, sheet


def _parse_headers(header_row):
    columns = []
    for index, raw in enumerate(header_row):
        if raw is None or str(raw).strip() == "":
            continue
        header = str(raw).strip()
        match = CAR_HEADER_RE.match(header)
        if match:
            columns.append(Column(index, header, car_key(match.group(1)),
                                  match.group(2).strip()))
        else:
            columns.append(Column(index, header, None, header))
    return columns


def _assign_circuit_channels(columns):
    for column in columns:
        if column.car is None:
            continue
        for channel, pattern in CIRCUIT_PATTERNS.items():
            match = pattern.search(column.param)
            if match:
                column.circuit_channel = channel
                column.circuit_index = match.group(1)
                break


def _resolve_roles_for_car(candidates, log, car):
    """Exact alias first, then keyword patterns; prefer a column that has data."""
    resolved = {}
    taken = set()
    for spec in ROLES:
        best = None
        # Alias before keyword, and a numeric role prefers a numeric column.  The
        # last two layers drop the numeric requirement so that a sensor which
        # names its role but reports a fault string ('Invalid' in place of a
        # temperature) is still recognised as that role -- otherwise the column
        # is never read and a dead probe cannot be reported at all.
        layers = (
            ("alias", spec.aliases, True),
            ("keyword", spec.keywords, True),
            ("alias (degraded)", spec.aliases, False),
            ("keyword (degraded)", spec.keywords, False),
        )
        for layer, patterns, strict in layers:
            pool = []
            for column in candidates:
                if column.index in taken:
                    continue
                param = column.param.lower()
                if any(neg in param for neg in spec.negative):
                    continue
                if strict and spec.numeric and column.nature == "categorical":
                    continue
                for rank, pattern in enumerate(patterns):
                    hit = (param == pattern) if layer == "alias" else (pattern in param)
                    if hit:
                        # Rank on the pattern's own priority first, then prefer a
                        # column that actually carries data over an empty one.
                        pool.append((rank, 0 if column.nature != "empty" else 1,
                                     column.index, column, layer))
                        break
            if pool:
                pool.sort(key=lambda item: item[:3])
                best = pool[0]
                break
        if best is None:
            continue
        _, _, _, column, layer = best
        column.role = spec.name
        column.matched_by = layer
        resolved[spec.name] = column
        taken.add(column.index)
        log.append("car {}: {} <- {!r} [{} match, nature={}]".format(
            car, spec.name, column.param, layer, column.nature))
    return resolved


def sniff(source, rows=SNIFF_ROWS, _escalated=False):
    """Phase 1.  Reads the header row plus `rows` data rows and nothing else."""
    workbook, sheet = _open_largest_sheet(source)
    try:
        stream = sheet.iter_rows(min_row=1, max_row=rows + 1, values_only=True)
        try:
            header_row = next(stream)
        except StopIteration:
            header_row = ()
        samples = [list(r) for r in stream]

        columns = _parse_headers(list(header_row))
        for column in columns:
            column.nature = classify_nature(
                [row[column.index] if column.index < len(row) else None
                 for row in samples]
            )
        _assign_circuit_channels(columns)

        cars = sorted({c.car for c in columns if c.car is not None})
        log = []
        roles = {}
        for car in cars:
            candidates = [c for c in columns if c.car == car]
            roles[car] = _resolve_roles_for_car(candidates, log, car)

        circuits = {}
        for car in cars:
            per_channel = {}
            for column in columns:
                if column.car != car or column.circuit_channel is None:
                    continue
                per_channel.setdefault(column.circuit_channel, {})[
                    column.circuit_index] = column
            circuits[car] = {ch: idx for ch, idx in per_channel.items()
                             if len(idx) >= 2}

        meta = {}
        for column in columns:
            if column.car is None:
                key = column.param.strip().lower()
                if key in ("car model", "train number", "time"):
                    meta[key] = column

        schema = Schema(
            sheet_title=sheet.title,
            n_rows=max((sheet.max_row or 1) - 1, 0),
            n_cols=sheet.max_column or len(columns),
            columns=columns,
            cars=cars,
            roles=roles,
            circuits=circuits,
            meta=meta,
            sniff_rows_used=rows,
            escalated=_escalated,
            log=log,
        )
    finally:
        workbook.close()

    # Bounded escalation: if a required role is unresolved for any car after the
    # 5-row sample, re-sniff once at 60 rows.  Never read the whole file.
    if not _escalated and rows < SNIFF_ROWS_ESCALATED and cars:
        unresolved = [car for car in cars
                      if any(r not in roles[car] for r in REQUIRED_ROLES)]
        if unresolved:
            deeper = sniff(source, rows=SNIFF_ROWS_ESCALATED, _escalated=True)
            deeper.log.insert(0, "ESCALATED to {} rows: required role unresolved "
                                 "for car(s) {} at {} rows".format(
                                     SNIFF_ROWS_ESCALATED,
                                     ", ".join(unresolved), rows))
            return deeper
    return schema


# --------------------------------------------------------------------------
# Phase 2 -- stream the sheet once, materialise only the columns with a role
# --------------------------------------------------------------------------

def _wanted_indices(schema):
    """Every column that earned a role, plus the per-circuit and meta columns.

    On the 483-column workbook this is a small fraction of the sheet.  Nothing
    else is materialised: pandas.read_excel(usecols=) is deliberately avoided
    because it still parses every cell into a frame.
    """
    wanted = set()
    for car in schema.cars:
        for column in schema.roles.get(car, {}).values():
            wanted.add(column.index)
        for channel in schema.circuits.get(car, {}).values():
            for column in channel.values():
                wanted.add(column.index)
    for column in schema.meta.values():
        wanted.add(column.index)
    return wanted


def materialise(source, schema):
    """Stream the chosen sheet once and keep only the wanted columns."""
    wanted = sorted(_wanted_indices(schema))
    workbook, sheet = _open_largest_sheet(source)
    data = {index: [] for index in wanted}
    n_rows = 0
    try:
        stream = sheet.iter_rows(values_only=True)
        next(stream, None)                      # header
        for row in stream:
            width = len(row)
            if all(cell is None for cell in row):
                continue
            for index in wanted:
                data[index].append(row[index] if index < width else None)
            n_rows += 1
    finally:
        workbook.close()
    return data, n_rows, len(wanted)


# --------------------------------------------------------------------------
# Regime filters
# --------------------------------------------------------------------------

def _min_rows(n_rows):
    return max(MIN_ROWS_ABS, math.ceil(MIN_ROWS_FRAC * n_rows))


DERATE_ACTIVE_TOKENS = ("halved", "shedding", "shed", "derate", "reduced", "limited")
DERATE_INACTIVE_TOKENS = ("normal", "no load", "none", "off", "inactive", "full")


def _derate_active(text):
    """True only when the value clearly says the unit is derating."""
    if text is None:
        return False
    if any(token in text for token in DERATE_INACTIVE_TOKENS):
        return False
    if any(token in text for token in DERATE_ACTIVE_TOKENS):
        return True
    number = to_number(text)
    if number is not None:
        return number != 0
    return False


def _cooling_states(values):
    """Pick the cooling vocabulary this file actually uses.

    Prefer a state that is both cooling AND automatic where one exists, because
    that is the regime the unit is being asked to hold a setpoint in.  Fall back
    to any state naming cooling.
    """
    seen = {v for v in values if v is not None}
    both = {v for v in seen if "cool" in v and "auto" in v}
    if both:
        return both, "cool+auto"
    any_cool = {v for v in seen if "cool" in v}
    return any_cool, "cool"


def _build_regime(data, schema, car, n_rows, log):
    """Rows where this car's unit is validly cooling.

    Every optional filter is applied only if it leaves at least
    max(20, 5% of rows); a filter that empties the data is the wrong filter for
    the schema, so it is dropped and logged.
    """
    roles = schema.roles.get(car, {})
    outcome_col = roles.get("outcome")
    reference_col = roles.get("reference")
    if outcome_col is None or reference_col is None:
        return [False] * n_rows, [], {}

    outcome = [to_number(v) for v in data[outcome_col.index]]
    reference = [to_number(v) for v in data[reference_col.index]]
    keep = [outcome[i] is not None and reference[i] is not None
            for i in range(n_rows)]
    applied = []
    detail = {"live_rows": sum(keep)}
    floor = _min_rows(n_rows)

    def try_filter(name, mask, note=""):
        candidate = [keep[i] and mask[i] for i in range(n_rows)]
        kept = sum(candidate)
        if kept >= floor:
            applied.append(name + (" (" + note + ")" if note else ""))
            detail[name] = kept
            return candidate
        log.append("car {}: SKIPPED {} -- would leave {} rows (floor {})".format(
            car, name, kept, floor))
        detail[name] = "skipped ({} rows)".format(kept)
        return keep

    valid_col = roles.get("information_valid")
    if valid_col is not None:
        texts = [to_text(v) for v in data[valid_col.index]]
        mask = [t is not None and "invalid" not in t and "valid" in t for t in texts]
        keep = try_filter("information_valid", mask)

    mode_col = roles.get("running_mode")
    if mode_col is not None:
        texts = [to_text(v) for v in data[mode_col.index]]
        states, kind = _cooling_states(texts)
        if states:
            mask = [t in states for t in texts]
            keep = try_filter("cooling_mode", mask, kind)
            detail["cooling_states"] = sorted(states)
        else:
            log.append("car {}: SKIPPED cooling_mode -- no state names a cooling "
                       "regime in this vocabulary".format(car))
            detail["cooling_mode"] = "skipped (no cooling state)"

    derate_col = roles.get("load_derate")
    if derate_col is not None:
        texts = [to_text(v) for v in data[derate_col.index]]
        mask = [not _derate_active(t) for t in texts]
        keep = try_filter("load_derate", mask)

    return keep, applied, detail


# --------------------------------------------------------------------------
# 4A -- cooling shortfall
# --------------------------------------------------------------------------

def _shortfall(data, schema, car, keep):
    """mean(outcome - reference) over the rows where the unit is validly cooling.

    Subtracting each car's own setpoint is essential, not cosmetic: in
    acv_case_01 cars 01-02 target 23.83 C while 03-08 target 24.16 C, so raw
    temperature promotes car 03 for being warm when it was merely told to be.
    """
    roles = schema.roles[car]
    outcome = [to_number(v) for v in data[roles["outcome"].index]]
    reference = [to_number(v) for v in data[roles["reference"].index]]
    deltas, outcomes, references = [], [], []
    for i, on in enumerate(keep):
        if not on:
            continue
        a, b = outcome[i], reference[i]
        if a is None or b is None:
            continue
        deltas.append(a - b)
        outcomes.append(a)
        references.append(b)
    return deltas, outcomes, references


# --------------------------------------------------------------------------
# 4F -- peak-load window (computed and displayed, never ranked on)
# --------------------------------------------------------------------------

def _peak_shortfall(data, schema, car, keep):
    """Shortfall restricted to the top decile of outdoor temperature.

    An undercharged circuit copes on a mild night and fails under thermal load,
    so this roughly doubles the separation margin -- but it weakens case_02, so
    it is corroboration only.
    """
    roles = schema.roles[car]
    outdoor_col = roles.get("outdoor_temperature")
    if outdoor_col is None:
        return None, 0
    outdoor = [to_number(v) for v in data[outdoor_col.index]]
    rows = [i for i, on in enumerate(keep) if on and outdoor[i] is not None]
    if len(rows) < MIN_ROWS_ABS:
        return None, 0
    ordered = sorted(rows, key=lambda i: outdoor[i], reverse=True)
    cut = max(MIN_ROWS_ABS, math.ceil(0.10 * len(ordered)))
    window = set(ordered[:cut])
    mask = [i in window for i in range(len(keep))]
    deltas, _, _ = _shortfall(data, schema, car, mask)
    return (nan_mean(deltas), len(deltas)) if deltas else (None, 0)


# --------------------------------------------------------------------------
# 4B -- circuit asymmetry
# --------------------------------------------------------------------------

def _circuit_series(data, column, keep):
    return [to_number(data[column.index][i]) for i, on in enumerate(keep) if on]


def _asymmetry(data, schema, car, keep, n_rows):
    """Relative imbalance between a car's own two refrigeration circuits.

    A refrigerant leak is localised to one circuit's pipework: it starves that
    circuit while the healthy one takes up the load.  A car's two circuits share
    the same cabin, setpoint, weather and duty cycle, so they are a far tighter
    control group than the other seven cars -- which is why WHOLE-CAR pressure,
    compression ratio and compressor duty are actively misleading here and are
    never used as ranking features.
    """
    channels = schema.circuits.get(car) or {}
    if not channels:
        return None, {}

    rows = [i for i, on in enumerate(keep) if on]
    if len(rows) < MIN_CIRCUIT_ROWS:
        rows = list(range(n_rows))
    mask = [False] * n_rows
    for i in rows:
        mask[i] = True

    per_channel = {}
    means = {}
    for channel, circuits in channels.items():
        indices = sorted(circuits)
        if len(indices) < 2:
            continue
        first = _circuit_series(data, circuits[indices[0]], mask)
        second = _circuit_series(data, circuits[indices[1]], mask)
        live_first = [v for v in first if v is not None]
        live_second = [v for v in second if v is not None]
        if len(live_first) < MIN_CIRCUIT_ROWS or len(live_second) < MIN_CIRCUIT_ROWS:
            continue
        mean_first, mean_second = nan_mean(live_first), nan_mean(live_second)
        means[channel] = (mean_first, mean_second)
        scale = (abs(mean_first) + abs(mean_second)) / 2.0
        if scale <= 1e-9:
            continue
        per_channel[channel] = abs(mean_first - mean_second) / scale

    # Compression ratio, when both pressures exist on both circuits.
    if "lp" in means and "hp" in means:
        ratios = []
        for side in (0, 1):
            low, high = means["lp"][side], means["hp"][side]
            ratios.append(high / low if low and abs(low) > 1e-9 else None)
        if all(r is not None for r in ratios):
            scale = (abs(ratios[0]) + abs(ratios[1])) / 2.0
            if scale > 1e-9:
                per_channel["ratio"] = abs(ratios[0] - ratios[1]) / scale
            means["ratio"] = (ratios[0], ratios[1])

    if not per_channel:
        return None, {}
    return nan_mean(list(per_channel.values())), {
        "channels": per_channel,
        "means": means,
        "rows_used": sum(mask),
    }


def _duty_by_quarter(data, schema, car, keep):
    """Per-circuit compressor duty by quarter of the recording.

    Display only.  On case_04 the faulty car reads Q1 0.93/0.34 ... Q4 0.67/0.95
    -- the starved circuit recovering, consistent with a recharge mid-window.
    This must not affect ranking.
    """
    channels = schema.circuits.get(car) or {}
    duty = channels.get("duty")
    if not duty or len(duty) < 2:
        return None
    rows = [i for i, on in enumerate(keep) if on]
    if len(rows) < 4 * MIN_CIRCUIT_ROWS:
        return None
    indices = sorted(duty)
    size = len(rows) // 4
    quarters = []
    for q in range(4):
        chunk = rows[q * size:(q + 1) * size] if q < 3 else rows[3 * size:]
        values = []
        for key in indices[:2]:
            series = data[duty[key].index]
            values.append(nan_mean([to_number(series[i]) for i in chunk]))
        quarters.append(tuple(values))
    return quarters


# --------------------------------------------------------------------------
# 4C -- cross-car z-score (corroboration, and the tie-breaker)
# --------------------------------------------------------------------------

def _zscore_anomaly(data, schema, n_rows, live_mask):
    """mean |z| per car, over parameters shared by every car.

    The outdoor-temperature role is excluded: weather is shared by the whole
    train, so comparing cars on it adds noise rather than signal.
    """
    cars = schema.cars
    by_param = {}
    for car in cars:
        for column in schema.columns:
            if column.car != car:
                continue
            # Outdoor temperature is shared by the whole train, so comparing
            # cars on the weather is noise.  The setpoint is excluded for the
            # same class of reason: it is a value the car was COMMANDED, not a
            # value it measured, so a car told to sit at a different temperature
            # is a configuration difference, not an anomaly.  That is the same
            # reasoning 4A uses when it subtracts each car's own setpoint.
            if column.role in ("outdoor_temperature", "reference"):
                continue
            if column.role is None and column.circuit_channel is None:
                continue
            if column.index not in data:
                continue
            by_param.setdefault(column.param.lower(), {})[car] = column

    usable = []
    for param, mapping in by_param.items():
        if len(mapping) != len(cars):
            continue
        # "numeric in >=50% of rows" is judged per car and only over the cars
        # that reported at all.  Pooling it across the whole train would let an
        # offline half of the trainset disqualify a parameter for the half that
        # is reporting -- which is exactly what happens on acv_case_04.
        reporting = 0
        for car, column in mapping.items():
            series = data[column.index]
            if not series:
                continue
            numeric = sum(1 for v in series if to_number(v) is not None)
            if numeric / len(series) >= 0.50:
                reporting += 1
        if reporting >= 3:
            usable.append((param, mapping))

    totals = {car: 0.0 for car in cars}
    counts = {car: 0 for car in cars}
    contributing = 0
    for _, mapping in usable:
        series = {car: [to_number(v) for v in data[column.index]]
                  for car, column in mapping.items()}
        before = sum(counts.values())
        for i in range(n_rows):
            present = [car for car in cars
                       if live_mask.get(car, [False] * n_rows)[i]
                       and series[car][i] is not None]
            if len(present) < 3:
                continue
            values = [series[car][i] for car in present]
            mu = sum(values) / len(values)
            var = sum((v - mu) ** 2 for v in values) / (len(values) - 1)
            sigma = math.sqrt(var)
            if sigma <= 1e-9:
                continue
            for car in present:
                totals[car] += abs((series[car][i] - mu) / sigma)
                counts[car] += 1
        # A parameter that is flat across the cars never separates them -- every
        # row was skipped for zero spread -- so it is not counted as a signal.
        if sum(counts.values()) > before:
            contributing += 1

    return ({car: (totals[car] / counts[car] if counts[car] else None)
             for car in cars},
            {car: counts[car] for car in cars},
            contributing)


# --------------------------------------------------------------------------
# 5 -- sensor-health alerts
# --------------------------------------------------------------------------

NUMERIC_ROLES = frozenset(spec.name for spec in ROLES if spec.numeric)


def _populations(data, schema, n_rows):
    """Populated fraction per (car, parameter) over the materialised columns.

    A measurement column counts as reporting only when it carries a number.
    Some units write a literal 'Invalid' into a temperature reading instead of
    leaving it blank -- that is a dead probe, not a populated cell, and treating
    it as populated would hide the fault.  Columns that are legitimately textual
    (running mode, information valid) keep the plain blank-or-not test.
    """
    by_param = {}
    for column in schema.columns:
        if column.car is None or column.index not in data:
            continue
        by_param.setdefault(column.param.lower(), []).append(column)

    populated = {}
    for param, columns in by_param.items():
        measurement = any(c.role in NUMERIC_ROLES or c.circuit_channel is not None
                          for c in columns)
        reported = {}
        for column in columns:
            series = data[column.index]
            if measurement:
                flags = [to_number(v) is not None for v in series]
            else:
                flags = [not is_null_like(v) for v in series]
            reported[column.car] = flags
        # Denominator is the window in which this parameter was recorded at all.
        # Rows where no car reported it are gaps in the recording, not evidence
        # against any one car, and counting them would drag every car below the
        # thresholds these rules depend on.
        window = 0
        for i in range(n_rows):
            if any(flags[i] for flags in reported.values()):
                window += 1
        for column in columns:
            live = sum(reported[column.car])
            populated[(column.car, param)] = (
                live / window if window else 0.0, live, column.param)
    return populated


def _alert(kind, severity, headline, seen, means, action, ranking_note, cars):
    return {
        "kind": kind,
        "severity": severity,           # 'fault' or 'informational'
        "headline": headline,
        "seen": seen,
        "means": means,
        "action": action,
        "ranking_note": ranking_note,
        "cars": cars,
    }


def _sensor_alerts(data, schema, n_rows, cars_info, regime_detail,
                   fleet_median=None, metric_key="shortfall"):
    alerts = []
    cars = schema.cars
    populated = _populations(data, schema, n_rows)
    params = sorted({key[1] for key in populated})

    def frac(car, param):
        return populated.get((car, param), (0.0, 0, param))[0]

    def label(param):
        for car in cars:
            if (car, param) in populated:
                return populated[(car, param)][2]
        return param

    # 1 -- CAR OFFLINE
    offline = [c for c in cars
               if params and all(frac(c, p) == 0.0 for p in params)]
    healthy = [c for c in cars
               if params and all(frac(c, p) >= 0.95 for p in params
                                 if any(frac(o, p) > 0 for o in cars))]
    if offline and len(offline) < len(cars) and healthy:
        listed = ", ".join(offline)
        alerts.append(_alert(
            "CAR OFFLINE", "fault",
            "CAR OFFLINE -- Cars " + listed,
            "{} sent no air-conditioning data at all for the whole recording, "
            "while cars {} reported normally throughout.".format(
                "These {} cars".format(len(offline)) if len(offline) > 1
                else "Car {}".format(offline[0]),
                ", ".join(c for c in cars if c not in offline)),
            "{} is not reaching the data recorder. This usually points to a "
            "communications gateway or a wiring fault between the two halves, "
            "rather than {} simultaneous sensor failures.".format(
                "Half the train" if len(offline) * 2 == len(cars)
                else "Part of the train", len(offline)),
            "Raise a comms fault for the {} unit before the next service. "
            "Do not clear these cars as healthy -- they have not been "
            "assessed.".format(
                "rear" if all(c > max(o for o in cars if o not in offline)
                              for c in offline) else "affected"),
            "these cars are placed last because they could not be measured, not "
            "because they are the least suspect.",
            offline))

    # 2 / 3 -- SENSOR DEAD and SENSOR INTERMITTENT, grouped by parameter so that
    # one finding produces one work order however many cars share it.
    def name_cars(listed):
        return "Car " + listed[0] if len(listed) == 1 else "Cars " + ", ".join(listed)

    for param in params:
        present = [c for c in cars if (c, param) in populated and c not in offline]
        if len(present) < 3:
            continue
        reporting = [c for c in present if frac(c, param) >= 0.95]
        if len(reporting) < 2:
            continue
        dead = sorted(c for c in present if frac(c, param) == 0.0)
        flaky = sorted(c for c in present if 0.05 <= frac(c, param) <= 0.90)
        reading = label(param).lower()

        if dead:
            alerts.append(_alert(
                "SENSOR DEAD", "fault",
                "SENSOR DEAD -- {}, {}".format(name_cars(dead), reading),
                "The {} is empty for the entire recording on {}, while "
                "the other {} car{} report it without gaps.".format(
                    reading, name_cars(dead).lower(), len(reporting),
                    "" if len(reporting) == 1 else "s"),
                "The probe or its wiring has most likely failed. The cooling "
                "itself may be fine -- but it is running blind, because it "
                "cannot see the value it is trying to control.",
                "Replace or re-terminate the probe on {}. Treat cooling "
                "performance as unverified until it reports again.".format(
                    name_cars(dead).lower()),
                "this reading is simply not used for the cars that lack it; "
                "they are still ranked on the readings they do report.",
                dead))

        if flaky:
            alerts.append(_alert(
                "SENSOR INTERMITTENT", "fault",
                "SENSOR INTERMITTENT -- {}, {}".format(name_cars(flaky), reading),
                "The {} drops in and out on {}, while the other cars "
                "report it throughout ({}).".format(
                    reading, name_cars(flaky).lower(),
                    ", ".join("car {} {:.0f}%".format(c, frac(c, param) * 100)
                              for c in flaky)),
                "That is usually a loose terminal, a chafed cable or a probe "
                "that is beginning to fail, rather than a fault in the cooling.",
                "Check the connection on {} at the next service.".format(
                    name_cars(flaky).lower()),
                "only the readings that were present are used, so the gaps do "
                "not distort those cars' scores.",
                flaky))

    # 4 -- SELF-CHECK DROPOUTS
    invalid_counts = {}
    for car in cars:
        column = schema.roles.get(car, {}).get("information_valid")
        if column is None or column.index not in data:
            continue
        texts = [to_text(v) for v in data[column.index]]
        invalid_counts[car] = sum(1 for t in texts
                                  if t is not None and "invalid" in t)
    if invalid_counts:
        flagged = [c for c, n in invalid_counts.items() if n > 0]
        clean = [c for c, n in invalid_counts.items() if n == 0]
        # The pattern is one car failing its self-check while every other car
        # reports valid on every row.  Several cars dropping out together is bus
        # noise affecting the whole train, not a finding against one carriage.
        if len(flagged) == 1 and len(clean) == len(invalid_counts) - 1:
            for car in flagged:
                alerts.append(_alert(
                    "SELF-CHECK DROPOUTS", "fault",
                    "SELF-CHECK DROPOUTS -- Car {}".format(car),
                    "Car {} marked its own data invalid on {} of {} readings. "
                    "Every other car reported valid data on every reading."
                    .format(car, invalid_counts[car], n_rows),
                    "The unit is intermittently failing its internal self-check, "
                    "or briefly dropping off the data bus. On its own this is "
                    "minor, but it is unusual for only one car to do it.",
                    "Note it on the work order for car {}. If that car is also "
                    "flagged for cooling shortfall, treat the two findings "
                    "together.".format(car),
                    "the invalid readings are excluded from the calculation, so "
                    "they do not distort car {}'s score.".format(car),
                    [car]))

    # 5 -- CONFIGURATION DIFFERENCE (header present for some cars only)
    header_params = {}
    for column in schema.columns:
        if column.car is None:
            continue
        header_params.setdefault(column.param.lower(), set()).add(column.car)
    partial = sorted(p for p, owners in header_params.items()
                     if 0 < len(owners) < len(cars))
    if partial:
        shown = ", ".join(partial[:4])
        alerts.append(_alert(
            "CONFIGURATION DIFFERENCE", "informational",
            "CONFIGURATION DIFFERENCE -- {} reading{}".format(
                len(partial), "" if len(partial) == 1 else "s"),
            "{} reading{} for some cars and not others ({}{}).".format(
                len(partial), " exists" if len(partial) == 1 else "s exist",
                shown, ", ..." if len(partial) > 4 else ""),
            "This is almost certainly how the trainset is built -- not every car "
            "carries the same equipment. It is not a fault.",
            "No action.",
            "unaffected. Only parameters shared by all cars are compared.",
            []))

    # 6 -- FUNCTION DISABLED (flat or zero for every car, all recording)
    for param in params:
        present = [c for c in cars if (c, param) in populated]
        if len(present) != len(cars):
            continue
        values = set()
        numeric_only = True
        for car in present:
            column = next((c for c in schema.columns
                           if c.car == car and c.param.lower() == param), None)
            if column is None or column.index not in data:
                numeric_only = False
                break
            for raw in data[column.index]:
                number = to_number(raw)
                if number is None:
                    if not is_null_like(raw):
                        numeric_only = False
                        break
                else:
                    values.add(number)
            if not numeric_only:
                break
        if numeric_only and len(values) == 1:
            only = values.pop()
            alerts.append(_alert(
                "FUNCTION DISABLED", "informational",
                "FUNCTION DISABLED -- {}".format(label(param).lower()),
                "The {} reads a flat {:g} on all eight cars for the whole "
                "recording.".format(label(param).lower(), only),
                "A value that never moves on any car is a function that is "
                "switched off for this service, not eight identical faults.",
                "No action.",
                "unaffected. A reading that is identical on every car cannot "
                "separate them.",
                []))

    # 7 -- SCHEMA MISMATCH
    skipped_everywhere = [
        name for name in ("cooling_mode", "information_valid", "load_derate")
        if cars and all(
            isinstance(regime_detail.get(c, {}).get(name), str)
            for c in cars)
        and any(name in regime_detail.get(c, {}) for c in cars)
    ]
    if skipped_everywhere:
        alerts.append(_alert(
            "SCHEMA MISMATCH", "fault",
            "SCHEMA MISMATCH -- {}".format(", ".join(skipped_everywhere)),
            "A regime filter ({}) matched almost none of the recording on any "
            "car.".format(", ".join(skipped_everywhere)),
            "This trainset labels its operating states differently from the rest "
            "of the fleet, so the usual filter does not fit. The readings "
            "themselves look fine.",
            "Confirm the state names for this train type with the supplier, then "
            "add them to the console's vocabulary.",
            "the filter was dropped rather than applied, so the cars are still "
            "compared on equal terms.",
            []))

    # 8 -- FILE NOT RANKABLE
    scoreable = [c for c, info in cars_info.items() if info["status"] == "scored"]
    if len(scoreable) < 2:
        alerts.append(_alert(
            "FILE NOT RANKABLE", "fault",
            "FILE NOT RANKABLE",
            "Only {} car{} produced usable cooling data in this file.".format(
                len(scoreable), "" if len(scoreable) == 1 else "s"),
            "With fewer than two measurable cars there is nothing to compare, so "
            "no carriage can be singled out.",
            "Check the data feed for this trainset before judging any carriage.",
            "all eight cars are still listed, in carriage order, so none is "
            "omitted -- but the order carries no diagnostic weight.",
            []))

    # FLEET-WIDE COOLING SHORTFALL -- a whole-train condition, not eight faults.
    scored_cars = [c for c, info in cars_info.items() if info["status"] == "scored"]
    if (metric_key == "shortfall" and fleet_median is not None
            and fleet_median > FLEET_SHORTFALL_ADVISORY_C and scored_cars):
        setpoints = [cars_info[c]["setpoint_mean"] for c in scored_cars
                     if cars_info[c]["setpoint_mean"] is not None]
        worst = max(scored_cars, key=lambda c: cars_info[c]["shortfall"])
        alerts.append(_alert(
            "FLEET-WIDE COOLING SHORTFALL", "advisory",
            "FLEET-WIDE COOLING SHORTFALL -- all {} measured cars".format(
                len(scored_cars)),
            "Every measured car on this trainset is running above its cooling "
            "target, by {:.2f} C for the middle car.".format(fleet_median),
            "The demanded cabin temperature on this trainset averages {:.1f} C. "
            "Units at capacity across the whole train point to ambient "
            "conditions or an unusually low setpoint rather than {} separate "
            "faults.".format(nan_mean(setpoints) if setpoints else float("nan"),
                             len(scored_cars)),
            "Review the setpoint for this trainset before raising any "
            "unit-level work. Car {} is still well outside its own fleet and "
            "should be inspected on its own merits.".format(worst),
            "unaffected. Cars are always compared against each other within the "
            "same trainset, so a shared condition cancels out.",
            sorted(scored_cars)))

    # PEAK-LOAD CORROBORATION UNAVAILABLE -- a missing corroborating channel,
    # stated plainly rather than silently producing nothing.
    starved = sorted(c for c in scored_cars
                     if not cars_info[c].get("peak_available"))
    if starved:
        whole_fleet = len(starved) == len(scored_cars)
        alerts.append(_alert(
            "PEAK-LOAD CORROBORATION UNAVAILABLE", "advisory",
            "PEAK-LOAD CORROBORATION UNAVAILABLE -- {}".format(
                "whole trainset" if whole_fleet
                else "Cars " + ", ".join(starved)),
            "The outdoor temperature sensor is not reporting on {} of the {} "
            "measured cars, so the hot-weather cross-check could not be run for "
            "{}.".format(len(starved), len(scored_cars),
                         "any of them" if whole_fleet else "those cars"),
            "An undercharged circuit copes on a mild night and fails under "
            "thermal load, so that check is how a marginal call is normally "
            "confirmed. Without it the cooling shortfall stands on its own.",
            "Report the outdoor probe on this trainset. Treat a narrow "
            "shortfall lead here as less well corroborated than usual.",
            "unaffected. The peak-load window is a cross-check only and is "
            "never part of the ranking.",
            starved))

    order = {"fault": 0, "advisory": 1, "informational": 2}
    alerts.sort(key=lambda a: (order[a["severity"]], a["kind"], a["headline"]))
    return alerts


# --------------------------------------------------------------------------
# Severity bands -- presentation only
# --------------------------------------------------------------------------

def _median(values):
    ordered = sorted(values)
    count = len(ordered)
    if not count:
        return None
    middle = count // 2
    if count % 2:
        return ordered[middle]
    return (ordered[middle - 1] + ordered[middle]) / 2.0


def fleet_stats(values, floor):
    """Median and floored MAD over the scored cars of one trainset."""
    live = [v for v in values if v is not None]
    if not live:
        return None, None
    med = _median(live)
    mad = max(_median([abs(v - med) for v in live]), floor)
    return med, mad


def robust_z(value, med, mad):
    if value is None or med is None or not mad:
        return None
    return (value - med) / mad


def severity_for(fleet_z, status):
    """Band from distance outside the car's own fleet.  Presentation only."""
    if status != "scored" or fleet_z is None:
        return "NO DATA", SEVERITY_ACTION["NO DATA"]
    for threshold, name in Z_BANDS:
        if fleet_z >= threshold:
            return name, SEVERITY_ACTION[name]
    return "NORMAL", SEVERITY_ACTION["NORMAL"]


# --------------------------------------------------------------------------
# Ranking
# --------------------------------------------------------------------------

STATUS_RANK = {"scored": 0, "unreliable": 1, "no_data": 2}


def _order_cars(cars_info, basis, log, z_params=0, n_rows=0):
    """Order every car.  No NaN ever reaches a sort key and no car is dropped."""
    cars = list(cars_info)

    def present(value):
        return 0 if value is not None else 1

    if basis == "circuit_asymmetry":
        keyed = sorted(cars, key=lambda c: (
            STATUS_RANK[cars_info[c]["status"]],
            present(cars_info[c]["asymmetry"]),
            -(cars_info[c]["asymmetry"] or 0.0),
            present(cars_info[c]["shortfall"]),
            -(cars_info[c]["shortfall"] or 0.0),
            c,
        ))
        return keyed

    keyed = sorted(cars, key=lambda c: (
        STATUS_RANK[cars_info[c]["status"]],
        present(cars_info[c]["shortfall"]),
        -(cars_info[c]["shortfall"] or 0.0),
        c,
    ))

    # 4D -- two cars whose shortfalls differ by less than TIE_BAND_C are
    # thermally indistinguishable.  Separate them with the cross-car z-score,
    # then with telemetry dropouts.  A single left-to-right pass of adjacent
    # swaps: no car can move past one that leads it by more than the band.
    index = 0
    while index < len(keyed) - 1:
        a, b = keyed[index], keyed[index + 1]
        info_a, info_b = cars_info[a], cars_info[b]
        if (info_a["status"] == info_b["status"] == "scored"
                and info_a["shortfall"] is not None
                and info_b["shortfall"] is not None
                and abs(info_a["shortfall"] - info_b["shortfall"]) < TIE_BAND_C):
            reason = None
            za, zb = info_a["anomaly"], info_b["anomaly"]
            # (a) the cross-car z-score ordering, but only where that score is
            # broad enough to be a call rather than a shortlist.
            if (z_params >= MIN_Z_PARAMS_FOR_TIEBREAK
                    and za is not None and zb is not None
                    and za > 0 and zb / za >= MARGIN_PROBABLE):
                reason = "cross-car z-score {:.3f} vs {:.3f} over {} parameters " \
                         "(margin {:.2f}x)".format(zb, za, z_params, zb / za)
            else:
                # (b) telemetry dropouts, where the difference is material.
                da, db = info_a["dropouts"], info_b["dropouts"]
                floor = MATERIAL_DROPOUT_FRAC * n_rows
                if db > da and (db - da) >= floor:
                    reason = "telemetry dropouts {} vs {} ({:.1%} of the " \
                             "recording apart)".format(db, da,
                                                       (db - da) / max(n_rows, 1))
            if reason:
                log.append(
                    "TIE-BREAK: cars {} and {} differ by {:.3f} C (< {:.2f}); "
                    "{} promoted on {}".format(
                        a, b, abs(info_a["shortfall"] - info_b["shortfall"]),
                        TIE_BAND_C, b, reason))
                keyed[index], keyed[index + 1] = b, a
                index += 2
                continue
        index += 1
    return keyed


# --------------------------------------------------------------------------
# RANKING BASES
#
# Two bases, never competing.  If the schema exposes per-circuit refrigeration
# data for at least 2 cars, rank by circuit asymmetry and break ties with the
# thermal shortfall.  Otherwise rank by shortfall.  `ranking_basis` is reported
# in every result so a silent switch is visible.
# --------------------------------------------------------------------------

FALLBACK_CARS = ["{:02d}".format(i) for i in range(1, 9)]


def _fallback_ranking(cars=None):
    return "|".join(cars if cars else FALLBACK_CARS)


def _nunique(values):
    return len({round(v, 6) for v in values if v is not None})


def rank_workbook(source, file_id="unknown.xlsx"):
    """Rank all 8 cars of one trainset.  The single scoring code path."""
    started = time.time()
    schema = sniff(source)
    data, n_rows, columns_read = materialise(source, schema)
    log = list(schema.log)
    floor = _min_rows(n_rows)

    cars_info = {}
    live_mask = {}
    regime_detail = {}

    for car in schema.cars:
        keep, applied, detail = _build_regime(data, schema, car, n_rows, log)
        live_mask[car] = keep
        regime_detail[car] = detail
        deltas, outcomes, references = _shortfall(data, schema, car, keep)
        kept = len(deltas)

        outcome_col = schema.roles.get(car, {}).get("outcome")
        dropouts = 0
        if outcome_col is not None and outcome_col.index in data:
            dropouts = sum(1 for v in data[outcome_col.index] if is_null_like(v))
        valid_col = schema.roles.get(car, {}).get("information_valid")
        if valid_col is not None and valid_col.index in data:
            dropouts += sum(1 for v in data[valid_col.index]
                            if (to_text(v) or "") == "invalid")

        # The checks of section 10, as status gates rather than crashes: a car
        # with no usable data gets an explicit status, never a NaN score.
        checks = []
        setpoint_mean = nan_mean(references)
        if detail.get("live_rows", 0) == 0 or kept == 0:
            status = "no_data"
            checks.append(("rows", False,
                           "car {}: 0/{} rows carried both a cabin temperature "
                           "and a setpoint".format(car, n_rows)))
        else:
            ok_rows = kept >= floor
            checks.append(("rows", ok_rows,
                           "car {}: {}/{} rows survived (floor {})".format(
                               car, kept, n_rows, floor)))
            ok_setpoint = setpoint_mean is not None and 5 < setpoint_mean < 40
            checks.append(("setpoint", ok_setpoint,
                           "car {}: setpoint mean {}".format(
                               car, "n/a" if setpoint_mean is None
                               else "{:.2f}".format(setpoint_mean))))
            ok_variety = _nunique(outcomes) > 3
            checks.append(("outcome_variety", ok_variety,
                           "car {}: outcome takes {} distinct values".format(
                               car, _nunique(outcomes))))
            status = "scored" if all(c[1] for c in checks) else "unreliable"

        shortfall = nan_mean(deltas) if status == "scored" else None
        peak, peak_n = _peak_shortfall(data, schema, car, keep) \
            if status == "scored" else (None, 0)
        asymmetry, asym_detail = _asymmetry(data, schema, car, keep, n_rows)

        cars_info[car] = {
            "car": car,
            "status": status,
            "shortfall": shortfall,
            "samples": kept,
            "live_rows": detail.get("live_rows", 0),
            "setpoint_mean": setpoint_mean,
            "outcome_mean": nan_mean(outcomes),
            "peak_shortfall": peak,
            "peak_samples": peak_n,
            "asymmetry": asymmetry,
            "asymmetry_detail": asym_detail,
            "duty_quarters": _duty_by_quarter(data, schema, car, keep),
            "dropouts": dropouts,
            "filters_applied": applied,
            "regime": detail,
            "checks": checks,
            "anomaly": None,
        }

    anomalies, anomaly_counts, z_params = _zscore_anomaly(
        data, schema, n_rows, live_mask)
    for car in schema.cars:
        cars_info[car]["anomaly"] = anomalies.get(car)
        cars_info[car]["anomaly_samples"] = anomaly_counts.get(car, 0)

    # Basis selection, gated behind the single constant.
    with_circuits = [c for c in schema.cars if cars_info[c]["asymmetry"] is not None]
    if RANKING_PREFER_ASYMMETRY and len(with_circuits) >= 2:
        basis = "circuit_asymmetry"
        basis_reason = ("per-circuit refrigeration data present for {} cars; "
                        "ranking on within-car circuit asymmetry, ties broken by "
                        "thermal shortfall".format(len(with_circuits)))
    else:
        basis = "cooling_shortfall"
        basis_reason = ("no per-circuit refrigeration data; ranking on cooling "
                        "shortfall against each car's own setpoint")

    tie_log = []
    if z_params < MIN_Z_PARAMS_FOR_TIEBREAK:
        tie_log.append(
            "TIE-BREAK: cross-car z-score rests on {} parameter(s), below the {} "
            "needed to be an independent signal; it is reported as corroboration "
            "only and does not reorder any car.".format(
                z_params, MIN_Z_PARAMS_FOR_TIEBREAK))
    ordered = _order_cars(cars_info, basis, tie_log, z_params, n_rows)
    log.extend(tie_log)

    # Hard invariants.
    sort_values = []
    for car in ordered:
        for key in ("shortfall", "asymmetry", "anomaly"):
            value = cars_info[car][key]
            if value is not None:
                sort_values.append(value)
    assert not any(isinstance(v, float) and math.isnan(v) for v in sort_values), \
        "NaN reached the sort key"
    assert len(ordered) == len(schema.cars), "a car was dropped from the ranking"
    assert sorted(ordered) == sorted(schema.cars), "ranking does not cover every car"

    # Severity is judged against this trainset's own fleet, on whichever metric
    # the file was actually ranked by, so the displayed number can never
    # contradict the displayed order.
    metric_key = "asymmetry" if basis == "circuit_asymmetry" else "shortfall"
    metric_floor = (MAD_FLOOR_RATIO if metric_key == "asymmetry"
                    else MAD_FLOOR_C)
    fleet_median, fleet_mad = fleet_stats(
        [cars_info[c][metric_key] for c in schema.cars
         if cars_info[c]["status"] == "scored"], metric_floor)

    for position, car in enumerate(ordered, start=1):
        info = cars_info[car]
        info["rank"] = position
        info["basis_metric"] = metric_key
        info["basis_value"] = info[metric_key]
        info["fleet_z"] = (robust_z(info[metric_key], fleet_median, fleet_mad)
                           if info["status"] == "scored" else None)
        band, action = severity_for(info["fleet_z"], info["status"])
        info["severity"] = band
        info["action"] = action
        info["peak_available"] = info["peak_shortfall"] is not None
        info["peak_agrees"] = (
            None if info["peak_shortfall"] is None or info["shortfall"] is None
            else (info["peak_shortfall"] >= info["shortfall"]))

    by_anomaly = sorted(
        [c for c in schema.cars if cars_info[c]["anomaly"] is not None],
        key=lambda c: -cars_info[c]["anomaly"])
    for position, car in enumerate(by_anomaly, start=1):
        cars_info[car]["anomaly_rank"] = position
    z_margin = None
    if len(by_anomaly) >= 2:
        top = cars_info[by_anomaly[0]]["anomaly"]
        second = cars_info[by_anomaly[1]]["anomaly"]
        if second and abs(second) > 1e-12:
            z_margin = top / second

    scored = [c for c in ordered if cars_info[c]["status"] == "scored"]
    lead = None
    if len(scored) >= 2:
        if basis == "circuit_asymmetry":
            first = cars_info[scored[0]]["asymmetry"]
            second = cars_info[scored[1]]["asymmetry"]
            lead = (first - second) if (first is not None and second is not None) else None
        else:
            first = cars_info[scored[0]]["shortfall"]
            second = cars_info[scored[1]]["shortfall"]
            lead = (first - second) if (first is not None and second is not None) else None

    # Rank 1 is always marked, whatever its band.  Severity answers "how bad";
    # rank answers "which one".  Without this, a trainset whose cars are all
    # within their own fleet shows eight unremarkable tiles and the controller
    # cannot see which car the model actually suspects.
    prime_ratio = None
    if len(scored) >= 2:
        first = cars_info[scored[0]][metric_key]
        second = cars_info[scored[1]][metric_key]
        if first is not None and second is not None and abs(second) > 1e-9:
            prime_ratio = first / second

    if not scored:
        prime_kind = "none"
    elif metric_key == "asymmetry":
        if prime_ratio is None or prime_ratio >= PRIME_RATIO_CLEAR:
            prime_kind = "clear"
        elif prime_ratio >= PRIME_RATIO_NARROW:
            prime_kind = "narrow"
        else:
            prime_kind = "tied"
    elif lead is None or lead >= CLEAR_LEAD_C:
        prime_kind = "clear"
    elif lead >= NARROW_LEAD_C:
        prime_kind = "narrow"
    else:
        prime_kind = "tied"

    prime_qualifier = {
        "clear": "clear lead",
        "narrow": "narrow lead, corroborate",
        "tied": "statistically tied with rank 2, treat as a shortlist",
        "none": "no carriage could be measured",
    }[prime_kind]

    alerts = _sensor_alerts(data, schema, n_rows, cars_info, regime_detail,
                            fleet_median, metric_key)

    meta_value = {}
    for key, column in schema.meta.items():
        series = data.get(column.index) or []
        values = [v for v in series if not is_null_like(v)]
        meta_value[key] = str(values[0]) if values else "-"

    return {
        "file_id": file_id,
        "ranked_cars": "|".join(ordered),
        "ranking_basis": basis,
        "basis_reason": basis_reason,
        "cars": [cars_info[c] for c in ordered],
        "cars_by_id": cars_info,
        "order": ordered,
        "alerts": alerts,
        "schema_log": log,
        "tie_log": tie_log,
        "columns_read": columns_read,
        "columns_total": schema.n_cols,
        "rows": n_rows,
        "sheet": schema.sheet_title,
        "escalated": schema.escalated,
        "car_model": meta_value.get("car model", "-"),
        "train_number": meta_value.get("train number", "-"),
        "z_margin": z_margin,
        "z_params": z_params,
        "lead": lead,
        "clear_lead": (prime_kind == "clear"),
        "prime_car": scored[0] if scored else (ordered[0] if ordered else None),
        "prime_kind": prime_kind,
        "prime_qualifier": prime_qualifier,
        "prime_ratio": prime_ratio,
        "basis_metric": metric_key,
        "fleet_median": fleet_median,
        "fleet_mad": fleet_mad,
        "peak_unavailable_cars": sorted(
            c for c in schema.cars
            if cars_info[c]["status"] == "scored"
            and not cars_info[c]["peak_available"]),
        "elapsed": time.time() - started,
        "ok": True,
        "error": None,
    }


def analyze_workbook(source, file_id="unknown.xlsx"):
    """rank_workbook, but a file that cannot be parsed still yields a ranking.

    An omitted car scores 0 while a last-placed car scores 0.125, so a fallback
    ordering is always better than no row.
    """
    try:
        return rank_workbook(source, file_id=file_id)
    except Exception as error:      # noqa: BLE001 -- a row must always be emitted
        cars = None
        try:
            cars = sniff(source).cars or None
        except Exception:           # noqa: BLE001
            cars = None
        ranked = _fallback_ranking(cars)
        return {
            "file_id": file_id,
            "ranked_cars": ranked,
            "ranking_basis": "fallback",
            "basis_reason": "file could not be analysed; carriage order emitted "
                            "so that no car is omitted",
            "cars": [], "cars_by_id": {}, "order": ranked.split("|"),
            "alerts": [_alert(
                "FILE NOT RANKABLE", "fault", "FILE NOT RANKABLE",
                "This trainset's file could not be read.",
                "The workbook is damaged, empty, or laid out differently from "
                "every other file in the batch.",
                "Re-export the recording for this trainset.",
                "all eight cars are listed in carriage order so none is omitted, "
                "but the order carries no diagnostic weight.", [])],
            "schema_log": ["FAILED: {}: {}".format(type(error).__name__, error)],
            "tie_log": [], "columns_read": 0, "columns_total": 0, "rows": 0,
            "sheet": "-", "escalated": False, "car_model": "-",
            "train_number": "-", "z_margin": None, "z_params": 0, "lead": None,
            "clear_lead": False, "prime_car": None, "prime_kind": "none",
            "prime_qualifier": "this trainset could not be read",
            "prime_ratio": None, "basis_metric": "shortfall",
            "fleet_median": None, "fleet_mad": None,
            "peak_unavailable_cars": [], "elapsed": 0.0, "ok": False,
            "error": "{}: {}".format(type(error).__name__, error),
        }


# --------------------------------------------------------------------------
# Submission
# --------------------------------------------------------------------------

SUBMISSION_COLUMNS = ("file_id", "ranked_cars")


def submission_rows(results):
    """Exactly two columns, built from the same string the engine returned."""
    return [{"file_id": r["file_id"], "ranked_cars": r["ranked_cars"]}
            for r in results]


def submission_csv(results):
    lines = [",".join(SUBMISSION_COLUMNS)]
    for row in submission_rows(results):
        lines.append("{},{}".format(row["file_id"], row["ranked_cars"]))
    return "\n".join(lines) + "\n"


def work_order_csv(results):
    """Plain-language depot work order.  No model internals."""
    header = ("trainset,carriage,priority,finding,recommended_action")
    lines = [header]
    for result in results:
        for info in result["cars"]:
            if info["severity"] in ("NORMAL",):
                continue
            if info["status"] == "scored":
                finding = ("cabin runs {:+.2f} C against its own setpoint"
                           .format(info["shortfall"]))
            else:
                finding = "no usable air-conditioning data was recorded"
            lines.append('{},{},{},"{}","{}"'.format(
                result["file_id"], info["car"], info["severity"],
                finding, info["action"]))
    return "\n".join(lines) + "\n"
