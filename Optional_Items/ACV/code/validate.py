"""Regression harness.  Exits non-zero on any failure.

Three layers, not one:

  1. Output      -- the exact ranked_cars string and ranking_basis per case.
  2. Intermediate -- per-car shortfall and asymmetry to 3 dp, eligible row count
                     per car, columns-read count, and the column chosen for each
                     role.  This is the layer that matters most.  A pipeline can
                     produce a correct final score from completely empty data:
                     if every shortfall came out NaN, pandas would sort the NaNs
                     into car order 01|02|...|08, and because case_04's true
                     faulty car happens to be 01 that would print score 1.000.
                     Ranks alone would pass that.  Row counts catch it instantly.
  3. Guard       -- deliberately broken copies of a real file, asserting that the
                     right alert fires and that no ranking is fabricated.

Ground truth is read here and nowhere else.  acv_core never sees it.
"""

from __future__ import annotations

import csv
import os
import shutil
import sys
import tempfile

import openpyxl

import acv_core as core

TRAIN_DIR = os.path.join("data", "Train")
TEST_FILE = os.path.join("data", "Test", "acv_test_case.xlsx")
LABELS = os.path.join("data", "Train_Labels.csv")


def load_truth():
    truth = {}
    with open(LABELS, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            truth[row["filename"].strip()] = core.car_key(row["faulty_car"])
    return truth


GOLDEN = {
    'acv_case_01.xlsx': {
        "ranked_cars": '01|02|03|04|07|08|05|06',
        "ranking_basis": 'cooling_shortfall',
        "basis_metric": 'shortfall',
        "prime_car": '01',
        "prime_kind": 'clear',
        "columns_read": 67, "columns_total": 67, "rows": 6999, "z_params": 1,
        "fleet_median": -0.205, "fleet_mad": 0.0596,
        "shortfall": {'01': 0.588, '02': 0.06, '03': -0.117, '04': -0.174, '05': -0.263, '06': -0.265, '07': -0.236, '08': -0.256},
        "asymmetry": {'01': None, '02': None, '03': None, '04': None, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': 13.3, '02': 4.44, '03': 1.47, '04': 0.52, '05': -0.98, '06': -1.02, '07': -0.52, '08': -0.86},
        "severity": {'01': 'CRITICAL', '02': 'HIGH', '03': 'NORMAL', '04': 'NORMAL', '05': 'NORMAL', '06': 'NORMAL', '07': 'NORMAL', '08': 'NORMAL'},
        "samples": {'01': 5794, '02': 5854, '03': 5842, '04': 5854, '05': 5854, '06': 5854, '07': 5854, '08': 5854},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'scored', '06': 'scored', '07': 'scored', '08': 'scored'},
        "roles_car01": {'heating_reference': 'ACV Control Temperature (Heating)', 'information_valid': 'ACV Information Valid', 'load_derate': 'Load Halved', 'outcome': 'Indoor Average Temperature', 'outdoor_temperature': 'Outdoor Average Temperature', 'reference': 'ACV Control Temperature (Cooling)', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Setting Mode'},
        "alerts": ['SELF-CHECK DROPOUTS'],
        "peak_unavailable_cars": [],
    },
    'acv_case_02.xlsx': {
        "ranked_cars": '02|03|07|08|06|01|04|05',
        "ranking_basis": 'cooling_shortfall',
        "basis_metric": 'shortfall',
        "prime_car": '02',
        "prime_kind": 'narrow',
        "columns_read": 67, "columns_total": 67, "rows": 9187, "z_params": 1,
        "fleet_median": 0.202, "fleet_mad": 0.05,
        "shortfall": {'01': 0.169, '02': 0.478, '03': 0.372, '04': 0.168, '05': 0.083, '06': 0.19, '07': 0.218, '08': 0.215},
        "asymmetry": {'01': None, '02': None, '03': None, '04': None, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': -0.67, '02': 5.5, '03': 3.4, '04': -0.68, '05': -2.38, '06': -0.25, '07': 0.32, '08': 0.25},
        "severity": {'01': 'NORMAL', '02': 'HIGH', '03': 'ELEVATED', '04': 'NORMAL', '05': 'NORMAL', '06': 'NORMAL', '07': 'NORMAL', '08': 'NORMAL'},
        "samples": {'01': 8291, '02': 8272, '03': 8290, '04': 8290, '05': 8291, '06': 8290, '07': 8290, '08': 8291},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'scored', '06': 'scored', '07': 'scored', '08': 'scored'},
        "roles_car01": {'heating_reference': 'ACV Control Temperature (Heating)', 'information_valid': 'ACV Information Valid', 'load_derate': 'Load Halved', 'outcome': 'Indoor Average Temperature', 'outdoor_temperature': 'Outdoor Average Temperature', 'reference': 'ACV Control Temperature (Cooling)', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Setting Mode'},
        "alerts": ['SELF-CHECK DROPOUTS'],
        "peak_unavailable_cars": [],
    },
    'acv_case_03.xlsx': {
        "ranked_cars": '03|02|01|07|04|08|05|06',
        "ranking_basis": 'cooling_shortfall',
        "basis_metric": 'shortfall',
        "prime_car": '03',
        "prime_kind": 'clear',
        "columns_read": 67, "columns_total": 67, "rows": 8310, "z_params": 1,
        "fleet_median": 0.319, "fleet_mad": 0.0837,
        "shortfall": {'01': 0.394, '02': 0.493, '03': 0.943, '04': 0.31, '05': 0.227, '06': 0.217, '07': 0.327, '08': 0.291},
        "asymmetry": {'01': None, '02': None, '03': None, '04': None, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': 0.9, '02': 2.08, '03': 7.46, '04': -0.1, '05': -1.1, '06': -1.21, '07': 0.1, '08': -0.33},
        "severity": {'01': 'NORMAL', '02': 'ELEVATED', '03': 'HIGH', '04': 'NORMAL', '05': 'NORMAL', '06': 'NORMAL', '07': 'NORMAL', '08': 'NORMAL'},
        "samples": {'01': 8265, '02': 8217, '03': 8216, '04': 8266, '05': 8266, '06': 8218, '07': 8218, '08': 7921},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'scored', '06': 'scored', '07': 'scored', '08': 'scored'},
        "roles_car01": {'heating_reference': 'ACV Control Temperature (Heating)', 'information_valid': 'ACV Information Valid', 'load_derate': 'Load Halved', 'outcome': 'Indoor Average Temperature', 'outdoor_temperature': 'Outdoor Average Temperature', 'reference': 'ACV Control Temperature (Cooling)', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Setting Mode'},
        "alerts": ['SELF-CHECK DROPOUTS'],
        "peak_unavailable_cars": [],
    },
    'acv_case_04.xlsx': {
        "ranked_cars": '01|04|03|02|05|06|07|08',
        "ranking_basis": 'circuit_asymmetry',
        "basis_metric": 'asymmetry',
        "prime_car": '01',
        "prime_kind": 'clear',
        "columns_read": 115, "columns_total": 483, "rows": 22262, "z_params": 9,
        "fleet_median": 0.032, "fleet_mad": 0.0107,
        "shortfall": {'01': -0.075, '02': -0.178, '03': -0.175, '04': 0.208, '05': None, '06': None, '07': None, '08': None},
        "asymmetry": {'01': 0.27, '02': 0.012, '03': 0.032, '04': 0.033, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': 22.32, '02': -1.93, '03': -0.07, '04': 0.07, '05': None, '06': None, '07': None, '08': None},
        "severity": {'01': 'CRITICAL', '02': 'NORMAL', '03': 'NORMAL', '04': 'NORMAL', '05': 'NO DATA', '06': 'NO DATA', '07': 'NO DATA', '08': 'NO DATA'},
        "samples": {'01': 20525, '02': 21172, '03': 21106, '04': 21196, '05': 0, '06': 0, '07': 0, '08': 0},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'no_data', '06': 'no_data', '07': 'no_data', '08': 'no_data'},
        "roles_car01": {'compressor_running': 'Compressor 1 Running', 'high_pressure': 'Refrigeration System 1 High Pressure Value', 'load_derate': 'Load Shedding', 'low_pressure': 'Refrigeration System 1 Low Pressure Value', 'outcome': 'Passenger Cabin Temperature Detected Value', 'outdoor_temperature': 'Fresh Air Temperature Detected Value', 'reference': 'Target Temperature Value', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Control Mode'},
        "alerts": ['CAR OFFLINE', 'CONFIGURATION DIFFERENCE'],
        "peak_unavailable_cars": [],
    },
    'acv_case_05.xlsx': {
        "ranked_cars": '04|02|07|01|06|03|08|05',
        "ranking_basis": 'cooling_shortfall',
        "basis_metric": 'shortfall',
        "prime_car": '04',
        "prime_kind": 'tied',
        "columns_read": 67, "columns_total": 67, "rows": 6972, "z_params": 1,
        "fleet_median": -0.924, "fleet_mad": 0.1022,
        "shortfall": {'01': -0.919, '02': -0.758, '03': -0.965, '04': -0.737, '05': -1.267, '06': -0.93, '07': -0.899, '08': -1.089},
        "asymmetry": {'01': None, '02': None, '03': None, '04': None, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': 0.06, '02': 1.63, '03': -0.39, '04': 1.83, '05': -3.36, '06': -0.06, '07': 0.25, '08': -1.61},
        "severity": {'01': 'NORMAL', '02': 'NORMAL', '03': 'NORMAL', '04': 'NORMAL', '05': 'NORMAL', '06': 'NORMAL', '07': 'NORMAL', '08': 'NORMAL'},
        "samples": {'01': 6211, '02': 6212, '03': 6212, '04': 6204, '05': 6212, '06': 6212, '07': 6212, '08': 6211},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'scored', '06': 'scored', '07': 'scored', '08': 'scored'},
        "roles_car01": {'heating_reference': 'ACV Control Temperature (Heating)', 'information_valid': 'ACV Information Valid', 'load_derate': 'Load Halved', 'outcome': 'Indoor Average Temperature', 'outdoor_temperature': 'Outside Temperature Sensor Reading', 'reference': 'ACV Control Temperature (Cooling)', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Setting Mode'},
        "alerts": ['FUNCTION DISABLED', 'PEAK-LOAD CORROBORATION UNAVAILABLE', 'SENSOR DEAD'],
        "peak_unavailable_cars": ['02', '03', '04', '05', '06', '07'],
    },
    'acv_case_06.xlsx': {
        "ranked_cars": '06|08|04|02|03|05|01|07',
        "ranking_basis": 'cooling_shortfall',
        "basis_metric": 'shortfall',
        "prime_car": '06',
        "prime_kind": 'clear',
        "columns_read": 67, "columns_total": 67, "rows": 3263, "z_params": 1,
        "fleet_median": 0.919, "fleet_mad": 0.0835,
        "shortfall": {'01': 0.711, '02': 0.929, '03': 0.908, '04': 0.966, '05': 0.903, '06': 2.238, '07': 0.653, '08': 1.039},
        "asymmetry": {'01': None, '02': None, '03': None, '04': None, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': -2.48, '02': 0.13, '03': -0.13, '04': 0.57, '05': -0.19, '06': 15.8, '07': -3.18, '08': 1.43},
        "severity": {'01': 'NORMAL', '02': 'NORMAL', '03': 'NORMAL', '04': 'NORMAL', '05': 'NORMAL', '06': 'CRITICAL', '07': 'NORMAL', '08': 'NORMAL'},
        "samples": {'01': 2623, '02': 2621, '03': 2622, '04': 2624, '05': 2621, '06': 2623, '07': 2622, '08': 2623},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'scored', '06': 'scored', '07': 'scored', '08': 'scored'},
        "roles_car01": {'heating_reference': 'ACV Control Temperature (Heating)', 'information_valid': 'ACV Information Valid', 'load_derate': 'Load Halved', 'outcome': 'Indoor Average Temperature', 'outdoor_temperature': 'Outside Temperature Sensor Reading', 'reference': 'ACV Control Temperature (Cooling)', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Setting Mode'},
        "alerts": ['FLEET-WIDE COOLING SHORTFALL', 'FUNCTION DISABLED', 'PEAK-LOAD CORROBORATION UNAVAILABLE', 'SENSOR DEAD'],
        "peak_unavailable_cars": ['02', '03', '04', '05', '06', '07'],
    },
    'acv_test_case.xlsx': {
        "ranked_cars": '01|03|04|07|08|06|02|05',
        "ranking_basis": 'cooling_shortfall',
        "basis_metric": 'shortfall',
        "prime_car": '01',
        "prime_kind": 'narrow',
        "columns_read": 67, "columns_total": 67, "rows": 9082, "z_params": 1,
        "fleet_median": -0.084, "fleet_mad": 0.05,
        "shortfall": {'01': 0.126, '02': -0.175, '03': -0.06, '04': -0.07, '05': -0.186, '06': -0.115, '07': -0.078, '08': -0.089},
        "asymmetry": {'01': None, '02': None, '03': None, '04': None, '05': None, '06': None, '07': None, '08': None},
        "fleet_z": {'01': 4.18, '02': -1.82, '03': 0.48, '04': 0.27, '05': -2.04, '06': -0.63, '07': 0.1, '08': -0.1},
        "severity": {'01': 'HIGH', '02': 'NORMAL', '03': 'NORMAL', '04': 'NORMAL', '05': 'NORMAL', '06': 'NORMAL', '07': 'NORMAL', '08': 'NORMAL'},
        "samples": {'01': 7798, '02': 7791, '03': 7792, '04': 7823, '05': 7841, '06': 7849, '07': 7859, '08': 7875},
        "status": {'01': 'scored', '02': 'scored', '03': 'scored', '04': 'scored', '05': 'scored', '06': 'scored', '07': 'scored', '08': 'scored'},
        "roles_car01": {'heating_reference': 'ACV Control Temperature (Heating)', 'information_valid': 'ACV Information Valid', 'load_derate': 'Load Halved', 'outcome': 'Indoor Average Temperature', 'outdoor_temperature': 'Outdoor Average Temperature', 'reference': 'ACV Control Temperature (Cooling)', 'running_mode': 'ACV Running Mode', 'setting_mode': 'ACV Setting Mode'},
        "alerts": [],
        "peak_unavailable_cars": [],
    },
}


FAILURES = []


def check(label, condition, detail=""):
    if condition:
        return True
    FAILURES.append("{}{}".format(label, (" -- " + detail) if detail else ""))
    return False


def approx(value, expected, places=3):
    if value is None or expected is None:
        return value is None and expected is None
    return round(float(value), places) == round(float(expected), places)


# --------------------------------------------------------------------------
# Layers 1 and 2
# --------------------------------------------------------------------------

def check_case(path, file_id, truth):
    gold = GOLDEN[file_id]
    result = core.rank_workbook(path, file_id)
    by_car = {info["car"]: info for info in result["cars"]}
    schema = core.sniff(path)

    # Layer 1 -- output.
    check("[{}] ranked_cars".format(file_id),
          result["ranked_cars"] == gold["ranked_cars"],
          "got {} expected {}".format(result["ranked_cars"], gold["ranked_cars"]))
    check("[{}] ranking_basis".format(file_id),
          result["ranking_basis"] == gold["ranking_basis"],
          "got {} expected {}".format(result["ranking_basis"], gold["ranking_basis"]))
    cars = result["ranked_cars"].split("|")
    check("[{}] all 8 cars present exactly once".format(file_id),
          len(cars) == 8 and len(set(cars)) == 8, str(cars))
    check("[{}] car ids are two-digit".format(file_id),
          all(len(c) == 2 and c.isdigit() for c in cars), str(cars))

    # Layer 2 -- intermediates.
    check("[{}] columns_read".format(file_id),
          result["columns_read"] == gold["columns_read"],
          "got {} expected {}".format(result["columns_read"], gold["columns_read"]))
    check("[{}] columns_total".format(file_id),
          result["columns_total"] == gold["columns_total"],
          "got {} expected {}".format(result["columns_total"], gold["columns_total"]))
    check("[{}] rows".format(file_id),
          result["rows"] == gold["rows"],
          "got {} expected {}".format(result["rows"], gold["rows"]))
    check("[{}] z_params".format(file_id),
          result["z_params"] == gold["z_params"],
          "got {} expected {}".format(result["z_params"], gold["z_params"]))

    # Display layer: the fleet-relative severity baseline and the rank-1 marker.
    check("[{}] basis_metric".format(file_id),
          result["basis_metric"] == gold["basis_metric"],
          "got {} expected {}".format(result["basis_metric"],
                                      gold["basis_metric"]))
    check("[{}] fleet_median".format(file_id),
          approx(result["fleet_median"], gold["fleet_median"]),
          "got {} expected {}".format(result["fleet_median"],
                                      gold["fleet_median"]))
    check("[{}] fleet_mad".format(file_id),
          approx(result["fleet_mad"], gold["fleet_mad"], 4),
          "got {} expected {}".format(result["fleet_mad"], gold["fleet_mad"]))
    check("[{}] prime_car".format(file_id),
          result["prime_car"] == gold["prime_car"],
          "got {} expected {}".format(result["prime_car"], gold["prime_car"]))
    check("[{}] prime_kind".format(file_id),
          result["prime_kind"] == gold["prime_kind"],
          "got {} expected {}".format(result["prime_kind"], gold["prime_kind"]))
    check("[{}] rank 1 is always marked".format(file_id),
          result["cars"][0]["rank"] == 1 if result["cars"] else True, "")
    check("[{}] peak_unavailable_cars".format(file_id),
          result["peak_unavailable_cars"] == gold["peak_unavailable_cars"],
          "got {} expected {}".format(result["peak_unavailable_cars"],
                                      gold["peak_unavailable_cars"]))

    for car in sorted(by_car):
        info = by_car[car]
        check("[{}] car {} shortfall".format(file_id, car),
              approx(info["shortfall"], gold["shortfall"][car]),
              "got {} expected {}".format(info["shortfall"], gold["shortfall"][car]))
        check("[{}] car {} asymmetry".format(file_id, car),
              approx(info["asymmetry"], gold["asymmetry"][car]),
              "got {} expected {}".format(info["asymmetry"], gold["asymmetry"][car]))
        check("[{}] car {} eligible rows".format(file_id, car),
              info["samples"] == gold["samples"][car],
              "got {} expected {}".format(info["samples"], gold["samples"][car]))
        check("[{}] car {} status".format(file_id, car),
              info["status"] == gold["status"][car],
              "got {} expected {}".format(info["status"], gold["status"][car]))
        check("[{}] car {} fleet z".format(file_id, car),
              approx(info["fleet_z"], gold["fleet_z"][car], 2),
              "got {} expected {}".format(info["fleet_z"],
                                          gold["fleet_z"][car]))
        check("[{}] car {} severity band".format(file_id, car),
              info["severity"] == gold["severity"][car],
              "got {} expected {}".format(info["severity"],
                                          gold["severity"][car]))

    resolved = {role: column.param
                for role, column in schema.roles["01"].items()}
    check("[{}] role columns for car 01".format(file_id),
          resolved == gold["roles_car01"],
          "got {} expected {}".format(resolved, gold["roles_car01"]))

    kinds = sorted({a["kind"] for a in result["alerts"]})
    check("[{}] alerts".format(file_id), kinds == gold["alerts"],
          "got {} expected {}".format(kinds, gold["alerts"]))

    # The invariants that must hold on every file, whatever the data.
    check("[{}] no NaN in any reported score".format(file_id),
          all(not (isinstance(info[key], float) and info[key] != info[key])
              for info in result["cars"]
              for key in ("shortfall", "asymmetry", "anomaly")
              if info[key] is not None), "")
    for info in result["cars"]:
        if info["status"] != "scored":
            continue
        floor = max(core.MIN_ROWS_ABS, 0.05 * result["rows"])
        check("[{}] car {}: rows survived".format(file_id, info["car"]),
              info["samples"] >= floor,
              "{}/{} rows survived (floor {:.0f})".format(
                  info["samples"], result["rows"], floor))
        check("[{}] car {}: setpoint plausible".format(file_id, info["car"]),
              info["setpoint_mean"] is not None and 5 < info["setpoint_mean"] < 40,
              "setpoint mean {}".format(info["setpoint_mean"]))

    rank = score = None
    if truth:
        order = result["ranked_cars"].split("|")
        rank = order.index(truth) + 1
        score = (len(order) - (rank - 1)) / len(order)
    return result, rank, score


# --------------------------------------------------------------------------
# Layer 3 -- guards
# --------------------------------------------------------------------------

def _rewrite(source, target, mutate):
    workbook = openpyxl.load_workbook(source)
    sheet = workbook.worksheets[0]
    header = [cell.value for cell in sheet[1]]
    mutate(sheet, header)
    workbook.save(target)


def guard_blank_car(workdir):
    """One car blanked entirely -> CAR OFFLINE, and that car is still ranked."""
    target = os.path.join(workdir, "blank_car.xlsx")

    def mutate(sheet, header):
        columns = [i + 1 for i, h in enumerate(header)
                   if h and str(h).strip().lower().startswith("car 05 -")]
        for row in range(2, sheet.max_row + 1):
            for column in columns:
                sheet.cell(row=row, column=column).value = None

    _rewrite(os.path.join(TRAIN_DIR, "acv_case_06.xlsx"), target, mutate)
    result = core.analyze_workbook(target, "blank_car.xlsx")
    cars = result["ranked_cars"].split("|")
    check("[guard blank car] all 8 cars still ranked",
          sorted(cars) == ["01", "02", "03", "04", "05", "06", "07", "08"],
          result["ranked_cars"])
    check("[guard blank car] car 05 reported as no_data",
          result["cars_by_id"].get("05", {}).get("status") == "no_data",
          str(result["cars_by_id"].get("05", {}).get("status")))
    check("[guard blank car] car 05 has no fabricated score",
          result["cars_by_id"].get("05", {}).get("shortfall") is None, "")
    check("[guard blank car] car 05 placed last",
          cars[-1] == "05", result["ranked_cars"])
    kinds = {a["kind"] for a in result["alerts"]}
    check("[guard blank car] CAR OFFLINE raised", "CAR OFFLINE" in kinds, str(kinds))
    return result


def guard_none_strings(workdir):
    """The literal string 'None' in place of blanks must not break scoring."""
    source = os.path.join(TRAIN_DIR, "acv_case_03.xlsx")
    target = os.path.join(workdir, "none_strings.xlsx")
    baseline = core.rank_workbook(source, "acv_case_03.xlsx")

    def mutate(sheet, header):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.value is None:
                    cell.value = "None"

    _rewrite(source, target, mutate)
    result = core.analyze_workbook(target, "none_strings.xlsx")
    check("[guard 'None' strings] ranking unchanged",
          result["ranked_cars"] == baseline["ranked_cars"],
          "got {} expected {}".format(result["ranked_cars"],
                                      baseline["ranked_cars"]))
    scored = sum(1 for info in result["cars"] if info["status"] == "scored")
    check("[guard 'None' strings] all 8 cars still scored", scored == 8,
          "{} cars scored".format(scored))
    check("[guard 'None' strings] no FILE NOT RANKABLE",
          "FILE NOT RANKABLE" not in {a["kind"] for a in result["alerts"]}, "")
    return result


def guard_unknown_modes(workdir):
    """A running-mode vocabulary naming no cooling state.

    The filter must be dropped rather than applied, and the file must still
    produce a full ranking off the remaining rows.
    """
    target = os.path.join(workdir, "unknown_modes.xlsx")

    def mutate(sheet, header):
        columns = [i + 1 for i, h in enumerate(header)
                   if h and "running mode" in str(h).strip().lower()]
        for row in range(2, sheet.max_row + 1):
            for column in columns:
                cell = sheet.cell(row=row, column=column)
                if cell.value is not None:
                    cell.value = "Betriebsart 7"

    _rewrite(os.path.join(TRAIN_DIR, "acv_case_06.xlsx"), target, mutate)
    result = core.analyze_workbook(target, "unknown_modes.xlsx")
    cars = result["ranked_cars"].split("|")
    check("[guard unknown modes] all 8 cars still ranked",
          sorted(cars) == ["01", "02", "03", "04", "05", "06", "07", "08"],
          result["ranked_cars"])
    check("[guard unknown modes] cooling filter was dropped, not applied",
          any("no state names a cooling regime" in line
              for line in result["schema_log"]),
          "; ".join(line for line in result["schema_log"] if "cooling" in line))
    scored = sum(1 for info in result["cars"] if info["status"] == "scored")
    check("[guard unknown modes] cars still scored from the remaining rows",
          scored == 8, "{} cars scored".format(scored))
    return result


def check_display_expectations(rows, test_result):
    """The display-layer expectations, asserted separately from the goldens."""
    by_file = {name: result for name, _, _, _, result in rows}
    by_file[test_result["file_id"]] = test_result

    def flagged(result):
        return [info["car"] for info in result["cars"]
                if info["severity"] not in ("NORMAL", "NO DATA")]

    # Fix 1 -- fleet-relative banding.
    six = by_file["acv_case_06.xlsx"]
    check("[fix1] case_06 flags exactly one car",
          flagged(six) == ["06"], str(flagged(six)))
    five = by_file["acv_case_05.xlsx"]
    check("[fix1] case_05 flags no car", flagged(five) == [], str(flagged(five)))
    for name, truth, _rank, _score, result in rows:
        ranked = sorted((info for info in result["cars"]
                         if info["fleet_z"] is not None),
                        key=lambda i: -i["fleet_z"])
        check("[fix1] {}: truth car {} has the highest fleet z".format(
            name, truth),
            ranked and ranked[0]["car"] == truth,
            "highest is {}".format(ranked[0]["car"] if ranked else "none"))

    # Severity must never enter a sort key: among the scored cars the ranking
    # must still be exactly the raw metric in descending order.
    for name, result in by_file.items():
        metric = result["basis_metric"]
        values = [info[metric] for info in result["cars"]
                  if info["status"] == "scored" and info[metric] is not None]
        check("[fix1] {}: ranking is still raw {} descending".format(
            name, metric),
            all(values[i] >= values[i + 1] for i in range(len(values) - 1)),
            str(values))

    # Fix 2 -- rank 1 is always marked, whatever its band.
    for name, result in by_file.items():
        check("[fix2] {}: a prime suspect is named".format(name),
              result["prime_car"] is not None, "")
    check("[fix2] case_05 reads as a shortlist, not a call",
          five["prime_kind"] == "tied", five["prime_kind"])

    # Fix 3 -- the headline metric matches the ranking basis.
    four = by_file["acv_case_04.xlsx"]
    check("[fix3] case_04 headline metric is circuit divergence",
          four["basis_metric"] == "asymmetry", four["basis_metric"])
    check("[fix3] case_04 car 01 bands CRITICAL on divergence",
          four["cars_by_id"]["01"]["severity"] == "CRITICAL",
          four["cars_by_id"]["01"]["severity"])

    # Fix 4 -- no-telemetry cars are an alert, and are still ranked.
    kinds = {a["kind"] for a in four["alerts"]}
    check("[fix4] case_04 raises CAR OFFLINE", "CAR OFFLINE" in kinds, str(kinds))
    offline = next((a for a in four["alerts"] if a["kind"] == "CAR OFFLINE"), None)
    check("[fix4] CAR OFFLINE names cars 05-08",
          offline is not None and offline["cars"] == ["05", "06", "07", "08"],
          str(offline["cars"]) if offline else "no alert")
    for car in ("05", "06", "07", "08"):
        check("[fix4] case_04 car {} still listed in ranked_cars".format(car),
              car in four["ranked_cars"].split("|"), four["ranked_cars"])
        check("[fix4] case_04 car {} carries no fabricated score".format(car),
              four["cars_by_id"][car]["shortfall"] is None, "")

    # Fix 5 -- advisories when a whole fleet is short, and when a corroborating
    # channel is unavailable.
    check("[fix5] case_06 raises FLEET-WIDE COOLING SHORTFALL",
          "FLEET-WIDE COOLING SHORTFALL" in {a["kind"] for a in six["alerts"]},
          str([a["kind"] for a in six["alerts"]]))
    for name in ("acv_case_05.xlsx", "acv_case_06.xlsx"):
        result = by_file[name]
        check("[fix5] {} reports the outdoor sensor as not reporting".format(name),
              "PEAK-LOAD CORROBORATION UNAVAILABLE"
              in {a["kind"] for a in result["alerts"]},
              str([a["kind"] for a in result["alerts"]]))
        check("[fix5] {}: prime suspect has no peak-load cross-check".format(name),
              result["prime_car"] in result["peak_unavailable_cars"],
              str(result["peak_unavailable_cars"]))
    for name in ("acv_case_01.xlsx", "acv_case_02.xlsx", "acv_case_03.xlsx"):
        result = by_file[name]
        check("[fix5] {} does not raise a spurious outdoor advisory".format(name),
              "PEAK-LOAD CORROBORATION UNAVAILABLE"
              not in {a["kind"] for a in result["alerts"]}, "")


def guard_submission_shape(results):
    text = core.submission_csv(results)
    header = text.splitlines()[0]
    check("[guard submission] header is exactly 'file_id,ranked_cars'",
          header == "file_id,ranked_cars", repr(header))
    for line in text.splitlines()[1:]:
        _file_id, _, ranked = line.partition(",")
        cars = ranked.split("|")
        check("[guard submission] 8 cars in every row",
              len(cars) == 8 and len(set(cars)) == 8, line)


# --------------------------------------------------------------------------

def main():
    truth = load_truth()
    rows = []
    results = []

    cases = sorted(name for name in os.listdir(TRAIN_DIR)
                   if name.endswith(".xlsx") and not name.startswith("~$"))
    for name in cases:
        result, rank, score = check_case(
            os.path.join(TRAIN_DIR, name), name, truth.get(name))
        results.append(result)
        rows.append((name, truth.get(name), rank, score, result))

    test_result, _, _ = check_case(TEST_FILE, os.path.basename(TEST_FILE), None)
    results.append(test_result)

    print("=" * 96)
    print("{:<20} {:<18} {:<26} {:>5} {:>7} {:>9} {:>9}".format(
        "file", "basis", "ranked_cars", "truth", "rank", "score", "margin"))
    print("-" * 96)
    for name, faulty, rank, score, result in rows:
        print("{:<20} {:<18} {:<26} {:>5} {:>7} {:>9.3f} {:>9}".format(
            name, result["ranking_basis"], result["ranked_cars"], faulty, rank,
            score, "-" if result["lead"] is None
            else "{:+.3f}".format(result["lead"])))
    scores = [score for _, _, _, score, _ in rows]
    mean = sum(scores) / len(scores)
    firsts = sum(1 for _, _, rank, _, _ in rows if rank == 1)
    print("-" * 96)
    print("{:<20} {:<18} {:<26} {:>5} {:>7} {:>9.3f}".format(
        "MEAN", "", "", "", "{}/{} at 1".format(firsts, len(rows)), mean))
    print("{:<20} {:<18} {:<26}".format(
        os.path.basename(TEST_FILE), test_result["ranking_basis"],
        test_result["ranked_cars"]))
    print("=" * 96)

    check_display_expectations(rows, test_result)

    workdir = tempfile.mkdtemp(prefix="acv_guard_")
    try:
        guard_blank_car(workdir)
        guard_none_strings(workdir)
        guard_unknown_modes(workdir)
        guard_submission_shape(results)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)

    check("mean score is 1.000", abs(mean - 1.0) < 1e-9, "{:.4f}".format(mean))
    check("six of six at rank 1", firsts == len(rows), "{} of {}".format(
        firsts, len(rows)))

    if FAILURES:
        print("\nFAILED ({} check{}):".format(
            len(FAILURES), "" if len(FAILURES) == 1 else "s"))
        for failure in FAILURES:
            print("  - {}".format(failure))
        return 1
    print("\nAll checks passed: output, intermediates and guards.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
